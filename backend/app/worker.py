"""Durable lease worker. Provider calls happen outside database transactions."""
import time, os, uuid, io, traceback, threading, zipfile
from sqlalchemy import select, update, or_, and_
from .core import Session, Record, Job, store, settings, uid, init, document_lock, execution_context
from . import core, dataset_models as routing
from . import providers, workflow
from .json_documents import json_document_text
from .relational import audit_as

MAX_ATTEMPTS=3
_last_reconcile=0.0
_last_cleanup=0.0

def claim_query(now, job_id=None, dialect='sqlite'):
    query=select(Job).where(or_(Job.status=='queued',and_(Job.status=='running',Job.lease_until<now)),Job.attempts<MAX_ATTEMPTS,Job.id==job_id if job_id else True).order_by(Job.created_at,Job.id).limit(1)
    # Oracle disallows FETCH FIRST together with FOR UPDATE. The conditional
    # UPDATE below still grants exactly one lease when workers race for a row.
    return query if dialect=='oracle' else query.with_for_update(skip_locked=True)

@audit_as('worker')
def reconcile_dependencies():
    # Jobs waiting for indexing survive API/worker restarts in SQL.
    with Session.begin() as session:
        for job in session.scalars(select(Job).where(Job.status=='blocked',Job.kind=='extract')):
            target=session.get(Record,job.target_id)
            if not target:job.status='cancelled';continue
            data=store.json(target.ref)
            documents=[session.get(Record,id) for id in data['document_ids']]
            if any(not d or d.status in ('failed','deleting') for d in documents):
                job.status='failed';job.error='Source indexing failed. Correct/retry the source job, then retry this extraction.'
                target.status='failed';target.error=job.error;workflow.fail(job.id,job.error)
            elif all(d.status=='ready' for d in documents):
                job.status='queued';target.status='queued';workflow.transition(job.id,'prepare','queued')

@audit_as('worker')
def claim(job_id=None):
    reconcile_dependencies()
    now=time.time();token=uid()
    with Session.begin() as s:
        for expired in s.scalars(select(Job).where(Job.status=='running',Job.lease_until<now,Job.attempts>=MAX_ATTEMPTS)):
            expired.status='failed';expired.error='Worker recovery attempts exhausted';expired.lease_until=0
            target=s.get(Record,expired.target_id)
            if target:target.status='failed';target.error=expired.error
            workflow.fail(expired.id,expired.error)
        candidate=s.scalar(claim_query(now,job_id,s.get_bind().dialect.name))
        if not candidate:return None
        result=s.execute(update(Job).where(Job.id==candidate.id,Job.attempts<MAX_ATTEMPTS,or_(Job.status=='queued',and_(Job.status=='running',Job.lease_until<now))).values(status='running',attempts=Job.attempts+1,lease_until=now+900,lease_token=token,updated_at=now))
        if result.rowcount!=1:return None
        return candidate.id,token

@audit_as('worker')
def checkpoint(jid,token,progress):
    with Session.begin() as s:
        j=s.get(Job,jid)
        if not j or j.status!='running' or j.lease_token!=token:raise InterruptedError('Job cancelled or lease lost')
        r=s.get(Record,j.target_id)
        if not r or r.status=='deleting':raise InterruptedError('Target deleted')
        j.progress=progress;j.lease_until=time.time()+900;j.updated_at=time.time()

def parse_document(name,raw,*,ocr_enabled=True,progress=None,cache_id=None):
    ext=name.rsplit('.',1)[-1].lower()
    if ext=='txt':return [{'page':1,'text':raw.decode('utf-8-sig')}]
    if ext=='json':return [{'page':1,'text':json_document_text(raw)}]
    if ext=='pdf':
        from pypdf import PdfReader
        reader=PdfReader(io.BytesIO(raw))
        if reader.is_encrypted:raise ValueError('Encrypted PDFs are unsupported. Upload an unlocked copy.')
        if len(reader.pages)>1000:raise ValueError('PDF exceeds 1000 page limit')
        pages=[];total=0
        for i,p in enumerate(reader.pages):
            text=p.extract_text() or ''
            # Image XObjects include scans even on pages with a selectable header.
            # Sparse text may also be a vector outline scan; try page OCR there.
            resources=p.get('/Resources',{})
            if hasattr(resources,'get_object'):resources=resources.get_object()
            objects=resources.get('/XObject',{})
            if hasattr(objects,'get_object'):objects=objects.get_object()
            def contains_image(objects,depth=0):
                if depth>8:raise ValueError('PDF image nesting exceeds the processing limit')
                for obj in objects.values():
                    obj=obj.get_object()
                    if obj.get('/Subtype')=='/Image':return True
                    if obj.get('/Subtype')=='/Form':
                        nested_resources=obj.get('/Resources',{})
                        if hasattr(nested_resources,'get_object'):nested_resources=nested_resources.get_object()
                        nested=nested_resources.get('/XObject',{})
                        if hasattr(nested,'get_object'):nested=nested.get_object()
                        if contains_image(nested,depth+1):return True
                return False
            visual=contains_image(objects) or bool(p.images.keys()) or len(text.strip())<20 or '\ufffd' in text
            value={'text':text,'text_origin':'native'}
            if visual and ocr_enabled:
                from . import ocr
                value=ocr.read_page(raw,i,text,progress=progress,cache_id=cache_id)
                text=value['text']
            elif visual:value['ocr_pending']=True
            total+=len(text)
            if total>5000000:raise ValueError('Document exceeds 5 million character text limit')
            pages.append({'page':i+1,**value})
        if ocr_enabled and not any(p['text'].strip() for p in pages):raise ValueError('No readable text was found after OCR. Upload a clearer scan or searchable PDF.')
        return pages
    if ext=='docx':
        from docx import Document
        try:archive=zipfile.ZipFile(io.BytesIO(raw))
        except zipfile.BadZipFile:raise ValueError('This DOCX file is invalid or damaged. Upload a valid DOCX copy.') from None
        with archive:
            entries=archive.infolist();names=[e.filename for e in entries]
            if len(names)!=len(set(names)) or any(n.startswith('/') or '..' in n.replace('\\','/').split('/') for n in names):raise ValueError('Unsafe DOCX archive entries')
            if sum(e.file_size for e in entries)>100*1024*1024 or any(e.file_size>25*1024*1024 for e in entries):raise ValueError('DOCX decompressed content exceeds safety limits')
            if '[Content_Types].xml' not in names or 'word/document.xml' not in names:raise ValueError('This file is not a valid DOCX document. Export it as DOCX and upload it again.')
        doc=Document(io.BytesIO(raw));lines=[p.text for p in doc.paragraphs]
        lines += [' | '.join(c.text for c in row.cells) for t in doc.tables for row in t.rows]
        text='\n'.join(lines)
        # DOCX image pagination cannot be inferred. Preserve the native text and
        # include every embedded raster image with an explicit image reference.
        if ocr_enabled:
            from . import ocr
            image_results=[]
            visible_images={}
            for part in doc.part.package.parts:
                # Package thumbnails and orphan media are not displayed content.
                if not str(part.partname).startswith('/word/') or not hasattr(part,'_element'):continue
                for relation_id in part._element.xpath('//@r:embed | //@r:id | //@r:link'):
                    relation=part.rels.get(relation_id)
                    if not relation:continue
                    if any(kind in relation.reltype for kind in ('/chart','/diagram','/oleObject')):
                        raise ValueError('DOCX contains a chart, diagram or embedded object. Export as PDF so its visual content can be read accurately.')
                    if not relation.reltype.endswith('/image'):continue
                    if relation.is_external:raise ValueError('DOCX has linked external images. Embed them or export as PDF; external image URLs are not fetched.')
                    visible_images[str(relation.target_part.partname)]=relation.target_part
            for n,part in enumerate(visible_images.values()):
                if part.content_type not in ('image/png','image/jpeg','image/bmp','image/tiff'):
                    raise ValueError('DOCX contains an unsupported image format. Export as PDF to OCR all visual content.')
                from PIL import Image
                image=Image.open(io.BytesIO(part.blob))
                if image.width*image.height>16000000:raise ValueError('DOCX image exceeds OCR pixel limit. Export as a smaller PDF.')
                image.load()
                if image.width<32 or image.height<32:continue
                out=io.BytesIO();image.convert('RGB').save(out,format='PNG')
                if progress:progress(n+1)
                result,meta=ocr.transcribe(out.getvalue(),ocr.available())
                image_results.append(meta|{'image':str(part.partname),'page':None})
                text+='\n[Embedded image '+str(part.partname)+'; page unknown]\n'+result
        if len(text)>5000000:raise ValueError('Document exceeds 5 million character text limit')
        return [{'page':None,'text':text,**({'text_origin':'native+ocr','ocr':{'model':ocr.config()['model'],'images':image_results,'verified':False,'notice':'DOCX image transcription; page location unavailable.'}} if ocr_enabled and image_results else {})}]
    raise ValueError('Unsupported document format')

def chunk_pages(did,pages,size,overlap):
    chunks=[]
    for page in pages:
        text=page['text']
        if len(chunks)+len(text)//(size-overlap)>5000:raise ValueError('Document exceeds 5000 chunk limit')
        for start in range(0,len(text),size-overlap):
            content=text[start:start+size]
            if not content.strip():continue
            ident=str(uuid.uuid5(uuid.NAMESPACE_URL,f'{did}:{page["page"]}:{start}:{providers.fingerprint()}:{content}'))
            chunks.append({'id':ident,'document_id':did,'page':page['page'],'offset':start,'text':content,**{k:page[k] for k in ('source_key','source_filename','source_page','text_origin','ocr') if k in page}})
    if settings()['embedding_provider']=='sentence_transformers' and settings()['search_provider']!='oci':
        model=providers.local_model('sentence-transformers/all-MiniLM-L6-v2');tokenizer=model.tokenizer;limit=model.max_seq_length-2;safe=[]
        for chunk in chunks:
            offsets=tokenizer(chunk['text'],add_special_tokens=False,return_offsets_mapping=True)['offset_mapping']
            for n in range(0,max(1,len(offsets)),limit):
                part=chunk.copy()
                if offsets:
                    begin=offsets[n][0];end=offsets[min(n+limit,len(offsets))-1][1];part['text']=chunk['text'][begin:end];part['offset']+=begin
                part['id']=str(uuid.uuid5(uuid.NAMESPACE_URL,chunk['id']+':'+str(n)));safe.append(part)
        chunks=safe
    elif settings()['embedding_provider']=='openai' and settings()['search_provider']!='oci':
        safe=[]
        for chunk in chunks:
            for offset,text in providers.split_utf8_bounded(chunk['text']):
                part=chunk.copy();part['text']=text;part['offset']+=offset
                part['id']=str(uuid.uuid5(uuid.NAMESPACE_URL,chunk['id']+':utf8:'+str(offset)));safe.append(part)
        chunks=safe
    return chunks

@audit_as('worker')
def process(jid,token):
    payload=store.json('jobs/'+jid+'.json');snapshot=payload.get('execution')
    if not snapshot:raise ValueError('Legacy job has no model snapshot. Configure its dataset and submit a new indexing or extraction job.')
    with execution_context(snapshot):
        with Session() as session:job=session.get(Job,jid)
        if job.kind=='index':
            with document_lock('dataset-'+snapshot['dataset_id']):
                dataset=routing.owned(snapshot['dataset_id'],job.owner,'knowledge_base')
                if routing.dataset_config(dataset).get('index_generation')!=snapshot['index_generation']:raise ValueError('Dataset embedding configuration changed. Submit a new reindex job.')
                _process(jid,token,snapshot)
        else:_process(jid,token,snapshot)

def _process(jid,token,snapshot):
    with Session() as s:
        j=s.get(Job,jid);record=s.get(Record,j.target_id)
        if not record:raise ValueError('Target no longer exists')
        kind=j.kind;did=record.id;name=record.name;owner=record.owner;kb=record.parent_id;ref=record.ref
    cfg=settings();checkpoint(jid,token,10)
    if kind=='index' and kb!=snapshot['dataset_id']:raise ValueError('Document dataset changed; submit a new indexing job.')
    if kind=='index':
        consent=store.json('jobs/'+jid+'.json').get('allow_external',False)
        if (cfg['embedding_provider'] in ('openai','registered') or cfg['search_provider']=='oci') and not consent:raise ValueError('External embedding/indexing requires explicit external data transmission consent')
        with Session.begin() as s:s.get(Record,did).status='processing'
        state=workflow.read(jid)
        # Retry an indexing failure from its saved extraction, using the same pinned configuration.
        reuse=state.get('current')=='index' and state.get('stages',{}).get('extract',{}).get('status')=='completed'
        if reuse:
            pages=store.json('documents/'+did+'/parsed.json')['pages']
            chunks=store.json('documents/'+did+'/chunks.json')['chunks']
        else:
            workflow.transition(jid,'extract','running')
            from . import merged_documents
            merged=merged_documents.manifest(did)
            def ocr_progress(page):
                checkpoint(jid,token,15)
                state=workflow.read(jid)
                state['stages']['extract'].update(detail='Reading/OCR page '+str(page),ocr_page=page)
                store.put_json('jobs/'+jid+'/workflow.json',state)
            if merged and (merged.get('ocr_pending') or any(p.get('ocr') for p in merged.get('pages',[]))):
                pages=merged_documents.read_sources(did,merged,progress=ocr_progress)
            else:pages=merged['pages'] if merged else parse_document(name,store.get(ref),progress=ocr_progress,cache_id=did)
            chunks=chunk_pages(did,pages,cfg['chunk_size'],cfg['chunk_overlap'])
        if not chunks:raise ValueError('Document has no usable text')
        if len(chunks)>5000:raise ValueError('Document exceeds 5000 chunk processing limit')
        from . import ocr
        store.put_json('documents/'+did+'/parsed.json',{'pages':pages,'ocr':ocr.summary(pages)})
        store.put_json('documents/'+did+'/chunks.json',{'chunks':chunks,'fingerprint':providers.fingerprint(),'chunk_size':cfg['chunk_size'],'chunk_overlap':cfg['chunk_overlap']})
        workflow.transition(jid,'extract','completed')
        workflow.transition(jid,'index','running')
        checkpoint(jid,token,35)
        vectors=[]
        if cfg['search_provider']=='oci':
            providers.delete_vectors(did)
            # Cleanup and crash recovery from this point belong to the NEW destination.
            store.put_json('documents/'+did+'/index-execution.json',snapshot)
            providers.oci_vector().upload_document(did,owner,kb,chunks,store,checkpoint=lambda:checkpoint(jid,token,50))
            dimensions=None
        else:
            for i in range(0,len(chunks),32):
                checkpoint(jid,token,35+int(45*i/len(chunks)))
                vectors.extend(providers.embed([c['text'] for c in chunks[i:i+32]]))
            checkpoint(jid,token,85)
            providers.delete_vectors(did)
            store.put_json('documents/'+did+'/index-execution.json',snapshot)
            providers.put_vectors(did,owner,kb,chunks,vectors)
            dimensions=len(vectors[0])
        try:checkpoint(jid,token,95)
        except InterruptedError:
            providers.delete_vectors(did)
            raise
        store.put_json('documents/'+did+'/index-execution.json',snapshot)
        store.put_json('documents/'+did+'/index.json',{'dataset_id':snapshot['dataset_id'],'index_generation':snapshot['index_generation'],'model_selection':routing.public_selection(snapshot),'search_provider':cfg['search_provider'],'fingerprint':providers.fingerprint(),'dimensions':dimensions,'chunk_count':len(chunks),'indexed_at':time.time(),'embedding_provider':cfg['embedding_provider'],'embedding_model':cfg['embedding_model'],'embedding_digest':cfg.get('_embedding_digest')})
        workflow.transition(jid,'index','completed')
        workflow.transition(jid,'ready','completed')
    elif kind=='extract':
        workflow.transition(jid,'prepare','running')
        data=store.json(ref)
        if cfg['model_provider'] in ('openai','oci','registered') and not data.get('allow_external'):raise ValueError('External extraction requires explicit external data transmission consent')
        sources=[]
        with Session() as s:
            for doc_id in data['document_ids']:
                d=s.get(Record,doc_id)
                from . import dataset_access
                if not d or not dataset_access.allowed(d,owner,s) or d.status!='ready' or d.parent_id!=snapshot['dataset_id']:raise ValueError('An extraction document is unavailable, outside the dataset, or not ready')
                if not routing.index_matches(d,snapshot):raise ValueError('Dataset index changed since extraction was submitted; run extraction again.')
                for c in store.json('documents/'+doc_id+'/chunks.json')['chunks']:sources.append({'document_id':doc_id,'document_name':d.name,**c})
        workflow.transition(jid,'prepare','completed')
        from . import extractor_agent
        with Session() as session:
            extraction_record=session.get(Record,did)
            documents=[session.get(Record,doc_id) for doc_id in data['document_ids']]
        data=extractor_agent.run(jid,extraction_record,data,snapshot,documents,sources,
            lambda:checkpoint(jid,token,20))
        result=data['result'];evidence=data['evidence']
        ocr_sources=[{'document_id':c['document_id'],'page':c['page'],'source_filename':c.get('source_filename',c['document_name']),'model':c['ocr']['model']} for c in sources if c.get('ocr')]
        ocr_sources=list({(c['document_id'],c['page']):c for c in ocr_sources}.values())
        data.update(result=result,evidence=evidence,ocr_sources=ocr_sources,evidence_notice='Review all values. Evidence contains only validated verbatim source quotes; fields without evidence are unverified.'+(' OCR was used: matching a transcription quote does not verify its accuracy against the original image. Check names, numbers and tables.' if ocr_sources else ''),sources=[{'document_id':c['document_id'],'document_name':c['document_name'],'chunk_id':c['id'],'page':c['page'],'excerpt':c['text'][:1000],**{k:c[k] for k in ('source_key','source_filename','source_page','text_origin') if k in c}} for c in sources],review_status='unreviewed',completed_at=time.time(),mock=cfg['model_provider']=='mock')
        store.put_json(ref,data)
        workflow.transition(jid,'ready','completed')
    else:raise ValueError('Unknown job type')
    with Session.begin() as s:
        j=s.get(Job,jid)
        if j.status!='running' or j.lease_token!=token:raise InterruptedError('Job cancelled')
        r=s.get(Record,did)
        if r:r.status='ready';r.error='';r.updated_at=time.time()
        j.status='completed';j.progress=100;j.error='';j.updated_at=time.time();j.lease_until=0

@audit_as('worker')
def run_once():
    from .queue_transport import get_transport
    global _last_reconcile,_last_cleanup
    transport=get_transport(settings()['queue_provider']);delivery=None;claimed=None
    now=time.monotonic()
    if now-_last_cleanup>=60:
        _last_cleanup=now
        from .temporary_files import purge_expired
        purge_expired()
    if now-_last_reconcile>=30:
        _last_reconcile=now;claimed=claim()
        # Confirmed Assistant chains continue even when the browser closes.
        from .agent_workspace import advance_workflows
        advance_workflows()
    if not claimed:
        delivery=transport.receive(timeout=0)
        claimed=claim(delivery.job_id) if delivery else None
    if delivery and not claimed:
        with Session() as session:
            existing=session.get(Job,delivery.job_id)
            if existing:transport.ack_terminal(delivery,existing.status)
    if not claimed:claimed=claim()
    if not claimed:return False
    jid,token=claimed
    stop=threading.Event()
    @audit_as('worker')
    def heartbeat():
        while not stop.wait(20):
            try:
                with Session.begin() as session:
                    job=session.get(Job,jid)
                    if not job or job.lease_token!=token or job.status!='running':return
                    job.lease_until=time.time()+900
                if delivery and delivery.job_id==jid:transport.renew(delivery)
            except Exception:pass
    thread=threading.Thread(target=heartbeat,daemon=True);thread.start()
    try:
        with Session() as session:target_id=session.get(Job,jid).target_id
        with document_lock(target_id):process(jid,token)
    except InterruptedError:pass
    except Exception as e:
        # Never log request bodies or provider secrets. Curated error classes only.
        message=str(e)[:400] if isinstance(e,(ValueError,providers.ProviderError)) else 'Processing failed; inspect service health and retry'
        with Session.begin() as s:
            j=s.get(Job,jid)
            if j and j.status=='running' and j.lease_token==token:
                j.status='failed';j.error=message;j.updated_at=time.time();j.lease_until=0
                r=s.get(Record,j.target_id)
                if r:r.status='failed';r.error=message
        workflow.fail(jid,message)
    finally:
        stop.set();thread.join(timeout=1)
        with Session() as session:finished=session.get(Job,jid)
        if finished and finished.status in ('completed','failed','cancelled'):
            from .agent_workspace import advance_workflows
            advance_workflows(job_id=jid)
        if delivery and delivery.job_id==jid:
            with Session() as session:finished=session.get(Job,jid)
            if finished:transport.ack_terminal(delivery,finished.status)
    return True

if __name__=='__main__':
    init()
    while True:
        if not run_once():time.sleep(2)
