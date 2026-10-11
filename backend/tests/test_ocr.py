"""OCR failures cannot silently lose pages or bypass merged-source processing."""
import io
import json
import pytest
from PIL import Image
from pypdf import PdfWriter
from app import ocr, worker, core, merged_documents
from test_dataset_models import system, dataset
from test_merged_documents import pdf


def scan(text='Scan',color='white'):
    image=Image.new('RGB',(240,160),color)
    from PIL import ImageDraw
    ImageDraw.Draw(image).text((20,20),text,fill='black')
    out=io.BytesIO();image.save(out,format='PDF');return out.getvalue()


@pytest.fixture
def vision(monkeypatch):
    requests=[]
    def call(cfg,path,payload=None,timeout=None):
        requests.append((path,payload))
        if path=='/api/show':return {'capabilities':['vision']}
        if path=='/api/tags':return {'models':[{'name':'qwen2.5vl:3b','digest':'a'*64}]}
        return {'done':True,'message':{'content':json.dumps({'text':'Alice Example\nInvoice total: 42.50','unreadable':False})}}
    monkeypatch.setattr(ocr,'call',call);return requests


def test_native_pdf_never_calls_ocr(vision):
    assert 'Alice' in worker.parse_document('native.pdf',pdf('Alice owns this fictional project.'))[0]['text']
    assert vision==[]


def test_scan_reads_page_and_retains_origin(vision):
    result=worker.parse_document('scan.pdf',scan())
    assert result[0]['text_origin']=='ocr' and result[0]['page']==1
    assert result[0]['ocr']['verified'] is False and result[0]['ocr']['coordinates'] is None
    request=next(p for path,p in vision if path=='/api/chat')
    assert request['messages'][1]['images'] and request['options']['temperature']==0.1
    assert request['format']=='json' and request['options']['repeat_penalty']==1.1
    assert 'image' not in ocr.summary(result)['pages'][0]


def test_blank_page_skips_model_but_blank_document_fails(vision):
    writer=PdfWriter();writer.add_blank_page(width=200,height=200);out=io.BytesIO();writer.write(out)
    with pytest.raises(ValueError,match='No readable text'):worker.parse_document('blank.pdf',out.getvalue())
    assert not vision


@pytest.mark.parametrize('response,message',[
    ({'done':False},'incomplete'),
    ({'done':True,'done_reason':'length'},'incomplete'),
    ({'done':True,'message':{'content':'not JSON'}},'invalid'),
    ({'done':True,'message':{'content':'{"text":"partial","unreadable":true}'}},'unreadable'),
    ({'done':True,'message':{'content':'{"text":"","unreadable":false}'}},'no text'),
    ({'done':True,'message':{'content':'{"text":"one","unreadable":false,"extra":1}'}},'invalid')])
def test_bad_vision_output_never_saved(monkeypatch,response,message):
    monkeypatch.setattr(ocr,'call',lambda *a,**k:response)
    out=io.BytesIO();Image.new('RGB',(240,160),'white').save(out,format='PNG')
    with pytest.raises(ocr.OcrError,match=message):ocr.transcribe(out.getvalue(),ocr.config()|{'digest':'a'})


def test_nonvision_model_rejected(monkeypatch):
    monkeypatch.setattr(ocr,'call',lambda *a,**k:{'capabilities':['completion']})
    with pytest.raises(ocr.OcrError,match='does not support images'):ocr.available()


def test_disabled_does_not_call_service(monkeypatch):
    monkeypatch.setenv('AEGIS_OCR_ENABLED','false')
    monkeypatch.setattr(ocr,'call',lambda *a,**k:pytest.fail('Disabled OCR contacted service'))
    with pytest.raises(ocr.OcrError,match='disabled'):ocr.available()


def test_visual_budget_fails_before_inference(monkeypatch):
    out=io.BytesIO();Image.new('RGB',(4000,4000),'white').save(out,format='PNG')
    monkeypatch.setattr(ocr,'call',lambda *a,**k:pytest.fail('Oversized page reached model'))
    with pytest.raises(ocr.OcrError,match='visual context budget'):ocr.transcribe(out.getvalue(),ocr.config())


def test_wide_word_image_padding_preserves_original_pixels(vision):
    import base64
    source=Image.new('RGB',(1100,260),(20,40,80));out=io.BytesIO();source.save(out,format='PNG')
    original=out.getvalue();ocr.transcribe(original,ocr.config()|{'digest':'a'})
    request=next(p for path,p in vision if path=='/api/chat')
    padded=Image.open(io.BytesIO(base64.b64decode(request['messages'][1]['images'][0])))
    assert padded.size==(1164,582)
    assert padded.crop((32,32,1132,292)).tobytes()==source.tobytes()
    assert padded.getpixel((0,0))==(255,255,255) and out.getvalue()==original


@pytest.mark.parametrize('key,value',[('AEGIS_OCR_CONTEXT','1024'),('AEGIS_OCR_DPI','900'),('AEGIS_OCR_MAX_TOKENS','20000'),('AEGIS_OCR_TIMEOUT','0'),('AEGIS_OCR_BATCH','0')])
def test_invalid_config_is_actionable(monkeypatch,key,value):
    monkeypatch.setenv(key,value)
    with pytest.raises(ocr.OcrError,match='configuration is invalid'):ocr.config()


def test_scan_merge_upload_does_not_run_model_and_worker_keeps_every_source(system,vision):
    d=dataset(system)
    parts=[('Alice.pdf',scan()),('Bob.pdf',scan('Bob'))]
    result=system.api.post('/api/documents/upload',data={'dataset_id':d['id'],'merge_files':'true'},files=[('files',(n,b,'application/pdf')) for n,b in parts])
    assert result.status_code==200,result.text
    assert not vision
    did=result.json()['documents'][0]['id']
    assert worker.run_once()
    doc=system.api.get('/api/documents/'+did).json()
    assert doc['status']=='ready' and doc['ocr']['ocr_page_count']==2
    saved=core.store.json('documents/'+did+'/parsed.json')['pages']
    assert [(p['page'],p['source_filename'],p['source_page']) for p in saved]==[(1,'Alice.pdf',1),(2,'Bob.pdf',1)]
    assert len([p for path,p in vision if path=='/api/chat'])==2
    assert core.store.get('documents/'+did+'/merged-sources/2/original')==parts[1][1]


def test_page_failure_cache_retry_only_complete_pages(system,vision,monkeypatch):
    raw=scan();did=core.uid();progress=[]
    first=ocr.read_page(raw,0,'',cache_id=did,progress=progress.append)
    second=ocr.read_page(raw,0,'',cache_id=did,progress=progress.append)
    assert first==second and progress==[1,1]
    assert len([p for path,p in vision if path=='/api/chat'])==1
    changed=scan('Different bytes')
    ocr.read_page(changed,0,'',cache_id=did)
    assert len([p for path,p in vision if path=='/api/chat'])==2


def test_mixed_native_and_scan_preserves_pages(vision):
    from pypdf import PdfReader
    writer=PdfWriter();writer.append(PdfReader(io.BytesIO(pdf('Native document header has known invoice details.'))));writer.append(PdfReader(io.BytesIO(scan())))
    out=io.BytesIO();writer.write(out)
    pages=worker.parse_document('mixed.pdf',out.getvalue())
    assert [p['text_origin'] for p in pages]==['native','ocr']
    assert [p['page'] for p in pages]==[1,2]


def test_docx_image_is_read_with_unknown_page(vision):
    from docx import Document
    doc=Document();doc.add_paragraph('Native owner is Bob Example.')
    out=io.BytesIO();Image.new('RGB',(240,160),'white').save(out,format='PNG');out.seek(0);doc.add_picture(out)
    data=io.BytesIO();doc.save(data)
    pages=worker.parse_document('mixed.docx',data.getvalue())
    assert pages[0]['page'] is None and 'Bob Example' in pages[0]['text'] and 'Alice Example' in pages[0]['text']
    assert pages[0]['ocr']['images'][0]['page'] is None


def test_document_chat_preserves_unverified_ocr_citation(system,vision):
    from test_dataset_models import conversation,message
    d=dataset(system)
    response=system.api.post('/api/documents/upload',data={'dataset_id':d['id']},files={'files':('scan.pdf',scan(),'application/pdf')})
    assert response.status_code==200,response.text
    did=response.json()['documents'][0]['id']
    assert worker.run_once()
    answer=message(system,conversation(system,d,[did]),text='What is the invoice total?')
    assert answer.status_code==200,answer.text
    citation=answer.json()['message']['citations'][0]
    assert citation['document_id']==did and citation['page']==1
    assert citation['text_origin']=='ocr' and citation['ocr']['verified'] is False
    assert citation['ocr']['coordinates'] is None and '42.50' in citation['text']
