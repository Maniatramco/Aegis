"""Cloud OCR preserves pinned credentials, external consent and page failure guards."""
import copy
import io
import json
import pytest
from PIL import Image
from ags_app import core, ocr, main, worker, dataset_models as routing, agent_consent
from ags_app.ags_models.connection import ModelConnection
from ags_app.ags_models import cloud_ocr
from test_dataset_models import system, dataset
from test_ocr import scan


def snapshot(provider='gemini', model=None):
    cfg = core.DEFAULTS | {'model_provider':'registered', 'model':model or ('gpt-4o' if provider=='openai' else 'gemini-3.5-flash-lite'),
        'context_limit':32768, 'max_output_tokens':2048, 'extra_parameters':{} if provider=='openai' else {'reasoning_effort':'minimal'},
        'model_connection':{'protocol':'openai-compatible','endpoint':'https://api.openai.com/v1' if provider=='openai' else 'https://generativelanguage.googleapis.com/v1beta/openai','auth_mode':'api_key'}}
    return {'settings':cfg,'provider_model':cfg['model'],'model_id':'cloud-model','model_version':1,
        'connection_profile_id':'cloud-profile','connection_profile_version':3,'secret_refs':{'connection':'pinned/cloud.enc'}}


@pytest.fixture(params=['gemini','openai'])
def cloud(monkeypatch, request):
    pinned = snapshot(request.param)
    def credentials(provider):
        if provider != 'connection':return None
        assert core._execution.get()['secret_refs']['connection'] == 'pinned/cloud.enc'
        return json.dumps({'api_key':'unit-test-credential'})
    monkeypatch.setattr(core, 'execution_secret', credentials)
    monkeypatch.setattr(ocr, 'call', lambda *a,**k:pytest.fail('Cloud OCR attempted local Ollama'))
    return pinned


def test_scan_uses_pinned_cloud_model_and_credentials_and_retains_provenance(cloud, monkeypatch):
    requests=[]
    def post(self, path, payload, headers=None):
        assert self.secret == {'api_key':'unit-test-credential'}
        requests.append((path,payload))
        return {'choices':[{'finish_reason':'stop','message':{'content':json.dumps({'text':'Cloud invoice total USD 29.50','unreadable':False})}}]}
    monkeypatch.setattr(ModelConnection,'post',post)
    with core.execution_context({'settings':core.DEFAULTS,'ocr_execution':cloud}):
        pages=worker.parse_document('scan.pdf',scan('Cloud invoice'))
    assert pages[0]['text_origin']=='ocr' and '29.50' in pages[0]['text']
    meta=pages[0]['ocr']
    assert meta['provider']==cloud_ocr.provider(cloud) and meta['model']==cloud['provider_model']
    assert meta['digest_type']=='configuration' and meta['verified'] is False and meta['coordinates'] is None
    path,payload=requests[0]
    assert path=='/chat/completions' and payload['model']==cloud['provider_model']
    assert payload['max_completion_tokens']==2048
    assert payload.get('reasoning_effort')==cloud['settings']['extra_parameters'].get('reasoning_effort')
    if meta['provider']=='openai':assert payload['messages'][1]['content'][1]['image_url']['detail']=='high'
    assert payload['response_format']['json_schema']['schema']==ocr.SCHEMA
    assert payload['messages'][1]['content'][1]['image_url']['url'].startswith('data:image/png;base64,')
    assert 'unit-test-credential' not in json.dumps(payload) and len(requests)==1
    assert ocr.summary(pages)['model']==cloud['provider_model']


@pytest.mark.parametrize('finish,content,reason',[
    ('length','{"text":"partial","unreadable":false}','incomplete'),
    ('content_filter',None,'incomplete'),
    ('stop','invalid JSON','invalid'),
    ('stop','{"text":"partial","unreadable":true}','unreadable'),
    ('stop','{"text":"","unreadable":false}','no text'),
    ('stop','{"text":"ok","unreadable":false,"extra":1}','invalid')])
def test_cloud_partial_refused_or_invalid_pages_are_never_accepted(cloud, monkeypatch, finish, content, reason):
    monkeypatch.setattr(ModelConnection,'post',lambda *a,**k:{'choices':[{'finish_reason':finish,'message':{'content':content}}]})
    image=io.BytesIO();Image.new('RGB',(240,160),'white').save(image,format='PNG')
    with core.execution_context({'settings':core.DEFAULTS,'ocr_execution':cloud}):
        with pytest.raises(ocr.OcrError,match=reason):ocr.transcribe(image.getvalue(),ocr.available())


def test_cloud_service_failure_does_not_leak_response_or_use_local_model(cloud, monkeypatch):
    def failure(*a,**k):raise RuntimeError('secret-provider-response')
    monkeypatch.setattr(ModelConnection,'post',failure)
    with core.execution_context({'settings':core.DEFAULTS,'ocr_execution':cloud}):
        with pytest.raises(ocr.OcrError,match=('OpenAI' if cloud_ocr.provider(cloud)=='openai' else 'Gemini')+' OCR request failed') as error:ocr.read_page(scan(),0,'')
    assert 'secret-provider-response' not in str(error.value)


def test_embedding_job_pins_dataset_extraction_registration_for_ocr(system, cloud, monkeypatch):
    s=system;item=dataset(s)
    record=routing.owned(item['id'],s.owner,'knowledge_base')
    real=routing.resolve_model
    monkeypatch.setattr(routing,'resolve_model',lambda r,cap,id=None:copy.deepcopy(cloud) if cap=='extraction' else real(r,cap,id))
    pinned=routing.execution_snapshot(record,'embedding')
    assert pinned['ocr_execution']['connection_profile_version']==3
    assert pinned['settings']['ocr_provider']==cloud_ocr.provider(cloud)
    cloud['connection_profile_version']=4
    assert pinned['ocr_execution']['connection_profile_version']==3


@pytest.mark.parametrize('provider',['gemini','openai'])
def test_cloud_ocr_alone_requires_consent_even_with_local_embeddings(provider):
    cfg=core.DEFAULTS|{'model_provider':'ollama','embedding_provider':'ollama','search_provider':'local','storage_provider':'local','ocr_provider':provider}
    with pytest.raises(main.HTTPException) as error:main.external_check(False,True,cfg)
    assert error.value.status_code==409
    notices=agent_consent.notices({'settings':cfg},'embedding',main)
    assert any(n['title']==('OpenAI' if provider=='openai' else 'Gemini')+' OCR' and 'images' in n['description'] for n in notices)


@pytest.mark.parametrize('provider',['gemini','openai'])
def test_cloud_ocr_configuration_version_changes_cache_identity(provider):
    first=snapshot(provider);second=copy.deepcopy(first);second['connection_profile_version']+=1
    limits={'enabled':True,'context':8192,'output':3000,'timeout':60,'dpi':180,'batch':128}
    adapter=cloud_ocr.adapter(provider)
    assert adapter.configure(first,limits)['digest']!=adapter.configure(second,limits)['digest']


@pytest.mark.parametrize('model',['gpt-4o','gpt-4o-mini','gpt-4o-2024-11-20','gpt-4o-mini-2024-07-18'])
def test_openai_ocr_supported_vision_aliases_and_snapshots(model):
    limits={'enabled':True,'context':8192,'output':3000,'timeout':1200,'dpi':180,'batch':128}
    cfg=cloud_ocr.adapter('openai').configure(snapshot('openai',model),limits)
    assert cfg['model']==model and cfg['output']==2048 and cfg['timeout']==core.DEFAULTS['timeout']
    assert cfg['context']==32768  # Registered cloud context, not local OCR's 8192.


@pytest.mark.parametrize('model,expected',[('gpt-4o',1105),('gpt-4o-mini',36835)])
def test_openai_landscape_page_uses_cloud_tile_budget(model,expected):
    adapter=cloud_ocr.adapter('openai')
    assert adapter.visual_tokens(2520,1530,model)==expected


@pytest.mark.parametrize('model',['gpt-4o','gpt-4o-mini'])
def test_large_openai_page_is_not_rejected_by_local_ocr_context_limit(monkeypatch,model):
    cloud=snapshot('openai',model);cloud['settings']['context_limit']=128000
    image=io.BytesIO();Image.new('RGB',(2520,1530),'white').save(image,format='PNG')
    sent=[]
    def post(self,path,payload,headers=None):
        sent.append(payload)
        return {'choices':[{'finish_reason':'stop','message':{'content':json.dumps({'text':'Synthetic landscape page','unreadable':False})}}]}
    monkeypatch.setattr(core,'execution_secret',lambda p:json.dumps({'api_key':'synthetic-only'}) if p=='connection' else None)
    monkeypatch.setattr(ModelConnection,'post',post)
    with core.execution_context({'settings':core.DEFAULTS,'ocr_execution':cloud}):
        text,meta=ocr.transcribe(image.getvalue(),ocr.available())
    assert text=='Synthetic landscape page' and len(sent)==1 and meta['model']==model


def test_openai_small_registered_budget_still_fails_before_inference(monkeypatch):
    cloud=snapshot('openai','gpt-4o-mini');cloud['settings']['context_limit']=8192
    image=io.BytesIO();Image.new('RGB',(2520,1530),'white').save(image,format='PNG')
    monkeypatch.setattr(ModelConnection,'post',lambda *a,**k:pytest.fail('Over-budget page was sent'))
    with core.execution_context({'settings':core.DEFAULTS,'ocr_execution':cloud}):
        with pytest.raises(ocr.OcrError,match='visual context budget'):
            ocr.transcribe(image.getvalue(),ocr.config())


@pytest.mark.parametrize('model',['text-embedding-3-small','gpt-3.5-turbo','gpt-4o-transcribe'])
def test_openai_ocr_rejects_embedding_and_nonvision_models_without_local_fallback(model, monkeypatch):
    monkeypatch.setattr(ocr,'call',lambda *a,**k:pytest.fail('Invalid cloud mapping contacted Ollama'))
    with core.execution_context({'settings':core.DEFAULTS,'ocr_execution':snapshot('openai',model)}):
        with pytest.raises(ocr.OcrError,match='configuration is invalid'):ocr.config()


def test_cloud_ocr_does_not_send_credentials_to_similarly_named_hosts():
    pinned=snapshot('openai')
    pinned['settings']['model_connection']['endpoint']='https://api.openai.com.example.org/v1'
    assert cloud_ocr.provider(pinned) is None


def test_ocr_consent_is_invalidated_by_changed_extraction_model_or_connection_version():
    cfg=core.DEFAULTS|{'model_provider':'ollama','embedding_provider':'ollama','search_provider':'local','storage_provider':'local',
        'ocr_provider':'openai','ocr_model':'gpt-4o','ocr_connection_profile_id':'ocr-profile','ocr_connection_profile_version':1,
        'ocr_connection':{'protocol':'openai-compatible','endpoint':'https://api.openai.com/v1'}}
    def key(values):return agent_consent.notices({'settings':values},'embedding',main)[0]['key']
    assert key(cfg)!=key(cfg|{'ocr_model':'gpt-4o-mini'})
    assert key(cfg)!=key(cfg|{'ocr_connection_profile_version':2})


def test_saved_invalid_extraction_mapping_never_selects_local_ocr(system,monkeypatch):
    item=dataset(system);record=routing.owned(item['id'],system.owner,'knowledge_base')
    def reject(*a,**k):raise main.HTTPException(409,'Selected extraction mapping is disabled.')
    monkeypatch.setattr(routing,'resolve_model',reject)
    with pytest.raises(main.HTTPException,match='disabled'):cloud_ocr.dataset_snapshot(record)


def test_cloud_scan_upload_requires_consent_and_worker_keeps_original_and_pinned_ocr(system,cloud,monkeypatch):
    item=dataset(system);real=routing.resolve_model
    monkeypatch.setattr(routing,'resolve_model',lambda r,cap,id=None:copy.deepcopy(cloud) if cap=='extraction' else real(r,cap,id))
    monkeypatch.setattr(ModelConnection,'post',lambda *a,**k:{'choices':[{'finish_reason':'stop','message':{'content':json.dumps({'text':'Synthetic invoice total USD 29.50','unreadable':False})}}]})
    raw=scan('Synthetic invoice')
    def upload(consent):return system.api.post('/api/documents/upload',files={'files':('cloud-scan.pdf',raw,'application/pdf')},data={'dataset_id':item['id'],'allow_external':str(consent).lower()})
    assert upload(False).status_code==409
    response=upload(True);assert response.status_code==200,response.text
    doc=response.json()['documents'][0]
    # Changing the currently resolved profile cannot change the queued image destination.
    cloud['connection_profile_version']=4
    assert worker.run_once()
    saved=system.api.get('/api/documents/'+doc['id']).json()
    assert saved['status']=='ready',saved
    record=routing.owned(doc['id'],system.owner,'document')
    assert core.store.get(record.ref)==raw
    parsed=core.store.json('documents/'+doc['id']+'/parsed.json')
    assert parsed['ocr']['provider']==cloud_ocr.provider(cloud)
    assert parsed['pages'][0]['text_origin']=='ocr' and '29.50' in parsed['pages'][0]['text']


def test_word_images_use_cloud_ocr_while_preserving_native_text(cloud,monkeypatch):
    from docx import Document
    image=io.BytesIO();Image.new('RGB',(240,160),'white').save(image,format='PNG');image.seek(0)
    document=Document();document.add_paragraph('Native invoice reference: NATIVE-17');document.add_picture(image)
    raw=io.BytesIO();document.save(raw)
    monkeypatch.setattr(ModelConnection,'post',lambda *a,**k:{'choices':[{'finish_reason':'stop','message':{'content':json.dumps({'text':'Embedded invoice total USD 29.50','unreadable':False})}}]})
    with core.execution_context({'settings':core.DEFAULTS,'ocr_execution':cloud}):
        pages=worker.parse_document('cloud-images.docx',raw.getvalue())
    assert 'NATIVE-17' in pages[0]['text'] and '29.50' in pages[0]['text']
    assert pages[0]['ocr']['provider']==cloud_ocr.provider(cloud)
    assert pages[0]['text_origin']=='native+ocr'
